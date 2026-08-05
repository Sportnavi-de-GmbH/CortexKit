"""Generate the non-technical Excel report package for the partner
recommendation agent.

Reads report-data.json (same directory) and writes eight .xlsx files.
Re-run after refreshing the data file:  python generate_reports.py

Design rules (per the reporting spec):
  - every workbook opens with a Summary sheet in plain language
  - human-readable column names, explanations, glossaries
  - status colours (green/amber/red), filters on every data table
  - charts where they help a non-technical reader
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).parent
DATA = json.loads((HERE / "report-data.json").read_text(encoding="utf-8"))

# ---------------------------------------------------------------- styling ---
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F3864")
SUB_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
STATUS_FILLS = {
    "good": PatternFill("solid", fgColor="C6EFCE"),
    "warn": PatternFill("solid", fgColor="FFEB9C"),
    "bad": PatternFill("solid", fgColor="FFC7CE"),
    "info": PatternFill("solid", fgColor="DDEBF7"),
}
STATUS_FONTS = {
    "good": Font(color="006100"),
    "warn": Font(color="9C6500"),
    "bad": Font(color="9C0006"),
    "info": Font(color="1F3864"),
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet_title(ws: Worksheet, title: str, subtitle: str, row: int = 1) -> int:
    ws.cell(row=row, column=1, value=title).font = TITLE_FONT
    ws.cell(row=row + 1, column=1, value=subtitle).font = SUB_FONT
    return row + 3


def put_table(
    ws: Worksheet,
    start_row: int,
    headers: list[str],
    rows: list[list],
    widths: list[int] | None = None,
    status_col: int | None = None,
    status_map: dict[str, str] | None = None,
    autofilter: bool = True,
) -> int:
    """Write a styled table; returns the row after the table."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row_vals in enumerate(rows, start_row + 1):
        for c, v in enumerate(row_vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.alignment = WRAP
        if status_col and status_map:
            key = str(row_vals[status_col - 1])
            kind = status_map.get(key)
            if kind:
                cell = ws.cell(row=r, column=status_col)
                cell.fill = STATUS_FILLS[kind]
                cell.font = STATUS_FONTS[kind]
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    if autofilter and rows:
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A{start_row}:{last_col}{start_row + len(rows)}"
    return start_row + len(rows) + 2


def put_glossary(ws: Worksheet, terms: list[tuple[str, str]]) -> None:
    row = sheet_title(
        ws,
        "Glossary",
        "Plain-language explanations of the technical terms used in this workbook.",
    )
    put_table(ws, row, ["Term", "What it means"], [list(t) for t in terms], [28, 100], autofilter=False)


STATUS3 = {"Healthy": "good", "Good": "good", "OK": "good", "Pass": "good",
           "Warning": "warn", "Watch": "warn", "Partial": "warn",
           "Critical": "bad", "Fail": "bad", "Blocked": "bad", "Info": "info"}


def save(wb: Workbook, name: str) -> None:
    wb.save(HERE / name)
    print(f"wrote {name}")


# =================================================== 1. Health Overview ====
def health_overview() -> None:
    d = DATA["health"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    row = sheet_title(
        ws,
        "AI Agent Health Overview — Partner Recommendation Agent",
        f"Executive dashboard · data collected {DATA['meta']['measured_on']} · "
        "for managers and product owners — no engineering knowledge required.",
    )
    row = put_table(
        ws,
        row,
        ["Metric", "Value", "Status", "What this means"],
        [[m["metric"], m["value"], m["status"], m["explanation"]] for m in d["summary"]],
        [30, 26, 12, 90],
        status_col=3,
        status_map=STATUS3,
    )
    ws.cell(row=row, column=1, value="How to read this sheet").font = Font(bold=True, size=12)
    for i, line in enumerate(d["how_to_read"], 1):
        ws.cell(row=row + i, column=1, value="• " + line).alignment = WRAP
    put_glossary(wb.create_sheet("Glossary"), [tuple(t) for t in d["glossary"]])
    save(wb, "AI-Agent-Health-Overview.xlsx")


# ============================================ 2. Workflow Performance ======
def workflow_performance() -> None:
    d = DATA["workflow"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = sheet_title(
        ws,
        "Agent Workflow Performance",
        "Every step the agent takes to answer one request, and how long each step takes.",
    )
    for i, line in enumerate(d["summary_lines"]):
        ws.cell(row=row + i, column=1, value="• " + line).alignment = WRAP
    ws2 = wb.create_sheet("Step Performance")
    row = sheet_title(
        ws2,
        "Performance per workflow step",
        "Times measured from real requests. 'Slowest seen' is the worst single execution observed.",
    )
    row = put_table(
        ws2,
        row,
        ["Agent Step", "What it does", "Average time", "Slowest seen",
         "Executions measured", "Failures", "Impact on speed", "Recommendation"],
        [[s["step"], s["description"], s["avg"], s["max"], s["n"], s["failures"],
          s["impact"], s["recommendation"]] for s in d["steps"]],
        [26, 46, 14, 14, 12, 10, 14, 46],
        status_col=7,
        status_map={"High": "bad", "Medium": "warn", "Low": "good", "Negligible": "good"},
    )
    # Chart: average ms per step (numeric column in a helper area)
    ws3 = wb.create_sheet("Chart Data")
    ws3.append(["Step", "Average ms"])
    for s in d["steps"]:
        if s.get("avg_ms") is not None:
            ws3.append([s["step"], s["avg_ms"]])
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Average time per step (milliseconds)"
    n = sum(1 for s in d["steps"] if s.get("avg_ms") is not None)
    data_ref = Reference(ws3, min_col=2, min_row=1, max_row=1 + n)
    cats = Reference(ws3, min_col=1, min_row=2, max_row=1 + n)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 10, 22
    ws2.add_chart(chart, f"A{row}")
    ws3.sheet_state = "hidden"
    put_glossary(wb.create_sheet("Glossary"), [tuple(t) for t in d["glossary"]])
    save(wb, "Agent-Workflow-Performance.xlsx")


# ================================================= 3. Errors & Issues ======
def errors_issues() -> None:
    d = DATA["errors"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = sheet_title(
        ws,
        "Sentry Errors and Issues",
        "Every problem the monitoring system has recorded for this agent, in plain language.",
    )
    for i, line in enumerate(d["summary_lines"]):
        ws.cell(row=row + i, column=1, value="• " + line).alignment = WRAP
    ws2 = wb.create_sheet("Issues")
    row = sheet_title(ws2, "Issue list", "One row per distinct problem. Use the filters to slice by severity or type.")
    put_table(
        ws2,
        row,
        ["Problem", "Type", "Severity", "Times seen", "First seen", "Last seen",
         "Status", "Root cause", "Recommended solution"],
        [[e["title"], e["class"], e["severity"], e["count"], e["first"], e["last"],
          e["status"], e["root_cause"], e["solution"]] for e in d["issues"]],
        [50, 14, 10, 10, 12, 12, 12, 45, 45],
        status_col=3,
        status_map={"fatal": "bad", "error": "warn", "warning": "info"},
    )
    ws3 = wb.create_sheet("Failure Classes")
    row = sheet_title(
        ws3,
        "The six failure classes",
        "Every captured problem is labelled with one of six classes so alerts stay understandable.",
    )
    put_table(
        ws3,
        row,
        ["Class", "Plain-English meaning", "Severity when it happens", "Verified reaching Sentry?"],
        [[c["name"], c["meaning"], c["severity"], c["verified"]] for c in d["classes"]],
        [16, 60, 22, 24],
        status_col=4,
        status_map={"Yes — VERIFIED 2026-07-24": "good"},
    )
    put_glossary(wb.create_sheet("Glossary"), [tuple(t) for t in d["glossary"]])
    save(wb, "Sentry-Errors-and-Issues.xlsx")


# ==================================================== 4. Latency ===========
def latency() -> None:
    d = DATA["latency"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = sheet_title(
        ws,
        "Latency Analysis — where the time goes",
        "How long users wait for an answer, and which component causes the wait.",
    )
    for i, line in enumerate(d["summary_lines"]):
        ws.cell(row=row + i, column=1, value="• " + line).alignment = WRAP
    ws2 = wb.create_sheet("Component Breakdown")
    row = sheet_title(ws2, "Time per component", "Averages from real measured requests.")
    row = put_table(
        ws2,
        row,
        ["Component", "Average time", "Share of total", "Assessment", "Note"],
        [[c["component"], c["avg"], c["share"], c["assessment"], c["note"]] for c in d["components"]],
        [34, 16, 14, 14, 70],
        status_col=4,
        status_map={"High": "bad", "Medium": "warn", "Low": "good"},
    )
    ws3 = wb.create_sheet("Chart Data")
    ws3.append(["Component", "Average ms"])
    for c in d["components"]:
        if c.get("avg_ms") is not None:
            ws3.append([c["component"], c["avg_ms"]])
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Average milliseconds per component"
    n = sum(1 for c in d["components"] if c.get("avg_ms") is not None)
    chart.add_data(Reference(ws3, min_col=2, min_row=1, max_row=1 + n), titles_from_data=True)
    chart.set_categories(Reference(ws3, min_col=1, min_row=2, max_row=1 + n))
    chart.height, chart.width = 9, 22
    ws2.add_chart(chart, f"A{row}")
    ws3.sheet_state = "hidden"
    ws4 = wb.create_sheet("Per Scenario")
    row = sheet_title(ws4, "End-to-end time per request type", "Wall-clock time a user would experience, single user.")
    put_table(
        ws4,
        row,
        ["Request type", "Example", "Total time", "Assessment"],
        [[s["scenario"], s["example"], s["total"], s["assessment"]] for s in d["scenarios"]],
        [24, 52, 14, 14],
        status_col=4,
        status_map={"Fast": "good", "Acceptable": "warn", "Slow": "bad"},
    )
    put_glossary(wb.create_sheet("Glossary"), [tuple(t) for t in d["glossary"]])
    save(wb, "Latency-Analysis.xlsx")


# ================================================ 5. Tokens & Cost =========
def tokens_cost() -> None:
    d = DATA["cost"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = sheet_title(
        ws,
        "Token and Cost Analysis",
        "What each AI request costs, measured from real traffic. "
        f"Prices: {d['pricing_note']}",
    )
    for i, line in enumerate(d["summary_lines"]):
        ws.cell(row=row + i, column=1, value="• " + line).alignment = WRAP
    ws2 = wb.create_sheet("Per Scenario")
    row = sheet_title(ws2, "Cost per request type", "Token counts are measured; costs follow from the prices above.")
    put_table(
        ws2,
        row,
        ["Scenario", "Model calls", "Input tokens", "Output tokens", "Cost (USD)", "Note"],
        [[s["scenario"], s["calls"], s["input"], s["output"], s["cost"], s["note"]] for s in d["scenarios"]],
        [34, 12, 14, 14, 12, 60],
    )
    ws3 = wb.create_sheet("Projections")
    row = sheet_title(ws3, "Monthly cost projections", "Straight-line projections from the measured average cost per request.")
    put_table(
        ws3,
        row,
        ["Traffic level", "Requests per day", "Cost per day", "Cost per month (30 days)"],
        [[p["level"], p["per_day"], p["day"], p["month"]] for p in d["projections"]],
        [22, 18, 14, 22],
    )
    ws4 = wb.create_sheet("Savings Levers")
    row = sheet_title(ws4, "How to reduce the cost", "Ordered by expected saving relative to effort.")
    put_table(
        ws4,
        row,
        ["Measure", "Expected saving", "Effort", "Explanation"],
        [[l["lever"], l["saving"], l["effort"], l["explanation"]] for l in d["levers"]],
        [36, 18, 10, 75],
        status_col=3,
        status_map={"Low": "good", "Medium": "warn", "High": "bad"},
    )
    put_glossary(wb.create_sheet("Glossary"), [tuple(t) for t in d["glossary"]])
    save(wb, "Token-and-Cost-Analysis.xlsx")


# ================================================ 6. Load Test =============
def load_test() -> None:
    d = DATA["load"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = sheet_title(
        ws,
        "Load Test Results",
        "How the agent behaves when several users ask at the same time.",
    )
    for i, line in enumerate(d["summary_lines"]):
        ws.cell(row=row + i, column=1, value="• " + line).alignment = WRAP
    ws2 = wb.create_sheet("Results")
    row = sheet_title(ws2, "Results per concurrency level", "Each level sends that many simultaneous user sessions.")
    row = put_table(
        ws2,
        row,
        ["Simultaneous users", "Requests", "Average response", "Slowest response",
         "Failed requests", "Result"],
        [[r["users"], r["requests"], r["avg"], r["max"], r["failed"], r["result"]] for r in d["levels"]],
        [20, 12, 18, 18, 16, 14],
        status_col=6,
        status_map=STATUS3,
    )
    ws3 = wb.create_sheet("Chart Data")
    ws3.append(["Users", "Average ms"])
    for r in d["levels"]:
        ws3.append([str(r["users"]), r["avg_ms"]])
    chart = BarChart()
    chart.title = "Average response time vs. simultaneous users"
    n = len(d["levels"])
    chart.add_data(Reference(ws3, min_col=2, min_row=1, max_row=1 + n), titles_from_data=True)
    chart.set_categories(Reference(ws3, min_col=1, min_row=2, max_row=1 + n))
    chart.height, chart.width = 9, 20
    ws2.add_chart(chart, f"A{row}")
    ws3.sheet_state = "hidden"
    ws4 = wb.create_sheet("Method")
    row = sheet_title(ws4, "How this was tested", "So the numbers can be reproduced.")
    for i, line in enumerate(d["method"]):
        ws4.cell(row=row + i, column=1, value="• " + line).alignment = WRAP
    put_glossary(wb.create_sheet("Glossary"), [tuple(t) for t in d["glossary"]])
    save(wb, "Load-Test-Results.xlsx")


# ============================================ 7. API Rate Limits ===========
def rate_limits() -> None:
    d = DATA["limits"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = sheet_title(
        ws,
        "API Rate Limit Analysis",
        "The external services this agent depends on, and how much headroom each one has.",
    )
    for i, line in enumerate(d["summary_lines"]):
        ws.cell(row=row + i, column=1, value="• " + line).alignment = WRAP
    ws2 = wb.create_sheet("Services")
    row = sheet_title(ws2, "Dependency limits", "'What happens at the limit' describes the user-visible effect.")
    put_table(
        ws2,
        row,
        ["Service", "Used for", "Limit", "Current usage", "Headroom",
         "What happens at the limit", "Recommendation"],
        [[s["service"], s["used_for"], s["limit"], s["usage"], s["headroom"],
          s["at_limit"], s["recommendation"]] for s in d["services"]],
        [24, 30, 26, 22, 12, 45, 45],
        status_col=5,
        status_map={"Large": "good", "Moderate": "warn", "Small": "bad", "Exhausted in bursts": "bad", "Unknown": "warn"},
    )
    put_glossary(wb.create_sheet("Glossary"), [tuple(t) for t in d["glossary"]])
    save(wb, "API-Rate-Limit-Analysis.xlsx")


# ======================================= 8. Optimization Priorities ========
def optimization_matrix() -> None:
    d = DATA["optimization"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    row = sheet_title(
        ws,
        "Optimization Priority Matrix",
        "What to fix first, scored by impact, saving, and difficulty. P1 = do now.",
    )
    for i, line in enumerate(d["summary_lines"]):
        ws.cell(row=row + i, column=1, value="• " + line).alignment = WRAP
    ws2 = wb.create_sheet("Priority Matrix")
    row = sheet_title(ws2, "All improvement items", "Sort or filter by any column.")
    put_table(
        ws2,
        row,
        ["Improvement", "Problem it solves", "Impact", "Cost saving", "Difficulty", "Priority", "Owner suggestion"],
        [[o["item"], o["problem"], o["impact"], o["saving"], o["difficulty"], o["priority"], o["owner"]] for o in d["items"]],
        [40, 55, 10, 12, 11, 10, 22],
        status_col=6,
        status_map={"P1": "bad", "P2": "warn", "P3": "info"},
    )
    put_glossary(wb.create_sheet("Glossary"), [tuple(t) for t in d["glossary"]])
    save(wb, "Optimization-Priority-Matrix.xlsx")


if __name__ == "__main__":
    health_overview()
    workflow_performance()
    errors_issues()
    latency()
    tokens_cost()
    load_test()
    rate_limits()
    optimization_matrix()
    print("done — 8 workbooks written to", HERE)
